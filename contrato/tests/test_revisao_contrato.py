#!/usr/bin/env python3
"""The revisao-contrato skill, proved without spending a single agent.

Run: python3 -m unittest discover -s contrato/tests   (needs `node`; python-docx for the docx case)

WHAT IS PROVED. That plan mode prints the plan line and dispatches no agent — `plan.mjs` runs the
real workflow script with `agent` stubbed to throw, so a spawn fails the run. That a simulated
review, driven by `sim.mjs` with a fake `agent`, pins a model on every spawn, stays under the cap,
hands every raw id to exactly one merged finding, sends only single-lens altas to the refuter, and
skips the conditional lenses when their condition is false. That the lens list in the script and
the sections of `lenses.md` are the same set. That the three Python helpers do what SKILL.md
tells the session they do, including the refusals.

WHAT IS NOT. Whether any lens finds anything in a real contract: that is the acceptance run,
which spends agents and is not a unit test.
"""

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.join(os.path.dirname(HERE), "skills", "revisao-contrato")
REF = os.path.join(SKILL, "reference")
SCRIPT = os.path.join(REF, "revisao-contrato.workflow.js")
NODE = shutil.which("node")

ARGS = {
    "contrato": "/t/contrato.md", "nossa_parte": "contratado", "autoria": "nossa",
    "acordo_previo": "/t/acordo.md", "data": "2026-09-21", "skill_dir": SKILL,
}


def plan(args):
    return subprocess.run([NODE, os.path.join(REF, "plan.mjs"), json.dumps(args)],
                          capture_output=True, text=True, timeout=60)


def sim(args, scenario=None):
    out = subprocess.run([NODE, os.path.join(HERE, "sim.mjs"), json.dumps(args), json.dumps(scenario or {})],
                         capture_output=True, text=True, timeout=60)
    if out.returncode != 0:
        raise AssertionError(out.stderr)
    return json.loads(out.stdout)


def helper(name, *argv, cwd=None):
    return subprocess.run([sys.executable, os.path.join(REF, name), *argv], capture_output=True, text=True, cwd=cwd, timeout=60)


@unittest.skipUnless(NODE, "node not installed")
class PlanLine(unittest.TestCase):
    def test_plan_prints_the_line_and_spawns_nothing(self):
        out = plan(ARGS)
        self.assertEqual(out.returncode, 0, out.stderr)
        first, saida, second = out.stdout.splitlines()[:3]
        self.assertEqual(saida, "saida: /t/revisao-contrato-2026-09-21")
        self.assertTrue(first.startswith("revisao-contrato · plano: ~22 agentes (teto 30;"), first)
        self.assertIn("opus:", first)
        self.assertIn("sonnet:", first)
        self.assertIn("janela de 5h", first)
        self.assertEqual(second, "agentes despachados neste plano: 0")

    def test_plan_mode_is_the_default_and_never_calls_agent(self):
        args = {k: v for k, v in ARGS.items()}
        d = sim(args, {"noAgent": True})  # agent stub throws when called
        self.assertIsNone(d["error"])
        self.assertEqual(d["calls"], [])
        self.assertEqual(d["result"]["agentes"]["previsto"], 22)
        self.assertFalse(d["result"]["perguntar"])

    def test_missing_args_are_named_and_ask(self):
        out = plan({"contrato": "/t/c.md", "skill_dir": SKILL})
        self.assertEqual(out.returncode, 3)
        self.assertRegex(out.stdout, r"FALTA: nossa_parte, autoria, acordo_previo, data")
        d = sim({"contrato": "/t/c.md", "autoria": "talvez"}, {"noAgent": True})
        self.assertTrue(d["result"]["perguntar"])
        self.assertIn("autoria (nossa|deles)", d["result"]["faltando"])

    def test_without_prior_agreement_fidelidade_leaves_the_plan(self):
        d = sim({**ARGS, "acordo_previo": "none"}, {"noAgent": True})
        self.assertEqual(d["result"]["agentes"]["previsto"], 21)

    def test_an_unknown_modo_stops_and_spawns_nothing(self):
        for modo in ("revisao", "apply", "Revisar"):
            d = sim({**ARGS, "modo": modo})
            self.assertIn(f'modo "{modo}" desconhecido', d["error"] or "", modo)
            self.assertEqual(d["calls"], [], modo)

    def test_review_refuses_to_start_with_missing_args(self):
        d = sim({**ARGS, "modo": "revisar", "data": "21/09"})
        self.assertIn("data (YYYY-MM-DD)", d["error"])
        self.assertEqual(d["calls"], [])


@unittest.skipUnless(NODE, "node not installed")
class SimulatedReview(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = sim({**ARGS, "modo": "revisar"}, {"extras": ["sucessao", "reversa"], "dropOne": True})

    def test_runs_to_the_report(self):
        self.assertIsNone(self.d["error"])
        self.assertEqual(self.d["calls"][-1]["label"], "relatorio")

    def test_every_spawn_pins_a_model_and_the_cap_holds(self):
        calls = self.d["calls"]
        self.assertTrue(all(c["model"] in ("opus", "sonnet") for c in calls))
        self.assertLessEqual(len(calls), 30)
        self.assertEqual(self.d["result"]["agentes"], len(calls))

    def test_models_follow_the_spec(self):
        by = {c["label"]: c for c in self.d["calls"]}
        self.assertEqual(by["lente:referencias-cruzadas"]["model"], "sonnet")
        for c in self.d["calls"]:
            if c["label"].startswith("consolidar"):
                self.assertEqual(c["model"], "sonnet")
        for label in ("scout", "critico", "relatorio", "lente:juridica-civil"):
            self.assertEqual(by[label]["model"], "opus")
        self.assertEqual(by["relatorio"]["effort"], "high")
        self.assertTrue(all(c["effort"] == "high" for c in self.d["calls"] if c["label"].startswith("verificar")))

    def test_effort_high_only_on_verifier_and_report(self):
        for c in self.d["calls"]:
            high = c["label"] == "relatorio" or c["label"].startswith("verificar")
            self.assertEqual(c["effort"], "high" if high else None, c["label"])

    def test_conditional_lenses(self):
        labels = [c["label"] for c in self.d["calls"]]
        self.assertNotIn("lente:trabalhista", labels)
        self.assertIn("lente:fidelidade", labels)
        self.assertIn("lente+:sucessao", labels)
        with_pf = sim({**ARGS, "modo": "revisar", "acordo_previo": "none"}, {"servicos_ou_pf": True})
        labels = [c["label"] for c in with_pf["calls"]]
        self.assertIn("lente:trabalhista", labels)
        self.assertNotIn("lente:fidelidade", labels)

    def test_only_fidelidade_gets_the_prior_agreement(self):
        for c in self.d["calls"]:
            if c["label"].startswith("lente"):
                self.assertEqual("/t/acordo.md" in c["prompt"], c["label"] == "lente:fidelidade", c["label"])

    def test_every_raw_id_in_exactly_one_merged_finding(self):
        # The fake consolidator drops one id; the check must pass it through raw.
        self.assertTrue(any("ausentes repassados" in l for l in self.d["result"]["checagem"]))
        c = self.d["result"]["contagens"]
        self.assertEqual(c["em_origem"], c["substantivos"])  # every raw id covered...
        self.assertEqual(c["origens"], c["substantivos"])    # ...and none twice
        for call in self.d["calls"]:
            if call["label"].startswith("consolidar"):
                items = json.loads(call["prompt"].split("RAW FINDINGS (JSON):", 1)[1])
                self.assertFalse([f for f in items if f["lente"] == "clareza"])  # clareza passes through

    def test_refuter_sees_only_single_lens_altas(self):
        for c in self.d["calls"]:
            if c["label"].startswith("verificar"):
                items = json.loads(c["prompt"].split("FINDINGS (JSON):", 1)[1])
                self.assertTrue(items)
                for f in items:
                    self.assertEqual(f["severidade"], "alta")
                    self.assertEqual(len({o.rsplit("-", 1)[0] for o in f["origem"]}), 1)

    def test_two_findings_of_one_lens_merged_are_still_single_lens(self):
        d = sim({**ARGS, "modo": "revisar"}, {"twiceFrom": "mesa"})
        self.assertIsNone(d["error"])
        refuted = [f for c in d["calls"] if c["label"].startswith("verificar")
                   for f in json.loads(c["prompt"].split("FINDINGS (JSON):", 1)[1])]
        merged = [f for f in refuted if f["clausula"] == "77.1"]
        self.assertEqual(len(merged), 1)
        self.assertEqual(sorted(merged[0]["origem"]), ["mesa-3", "mesa-4"])
        self.assertEqual(merged[0]["convergencia"], 1)

    def test_an_extra_lens_never_takes_an_active_lens_key(self):
        d = sim({**ARGS, "modo": "revisar"}, {"extras": ["Tributária", "tributaria"]})
        self.assertIsNone(d["error"])
        extra = sorted(c["label"] for c in d["calls"] if c["label"].startswith("lente+:"))
        self.assertEqual(extra, ["lente+:tributaria-extra", "lente+:tributaria-extra2"])
        brutos = [re.search(r"/brutos/([^ ]+)\.json", c["prompt"]).group(1)
                  for c in d["calls"] if c["label"].startswith("lente")]
        self.assertEqual(len(brutos), len(set(brutos)))
        self.assertEqual(d["result"]["contagens"]["em_origem"], d["result"]["contagens"]["substantivos"])

    def test_a_leading_clause_number_wins_over_the_parties_words(self):
        d = sim({**ARGS, "modo": "revisar"}, {"anchorFrom": "mesa", "anchor": "7.3 (partes relacionadas)"})
        self.assertIsNone(d["error"])
        jobs = [c for c in d["calls"] if c["label"].startswith("consolidar")]
        owner = [c["label"] for c in jobs if "7.3 (partes relacionadas)" in c["prompt"]]
        self.assertEqual(len(owner), 1)
        names = owner[0].split(":", 1)[1].split("+")
        self.assertIn("c7", names)
        self.assertNotIn("partes", names)

    def test_lens_prompts_carry_no_party_facts_of_their_own(self):
        # The script is generic: the only paths in a lens prompt are those that came in by args.
        lens = next(c for c in self.d["calls"] if c["label"] == "lente:juridica-civil")["prompt"]
        self.assertIn('good for the party "contratado"', lens)
        self.assertIn(os.path.join(REF, "lenses.md"), lens)

    def test_apply_mode_one_agent_per_clause_group(self):
        d = sim({**ARGS, "modo": "aplicar", "saida": "/t/rev", "aplicar": ["c4-01", "c4-02", "anexo-iv-01"]})
        self.assertIsNone(d["error"])
        self.assertEqual(sorted(c["label"] for c in d["calls"]), ["aplicar:anexo-iv", "aplicar:c4"])
        self.assertTrue(all(c["model"] == "opus" for c in d["calls"]))
        d = sim({**ARGS, "modo": "aplicar", "saida": "/t/rev", "aplicar": "todas"})
        self.assertIn("explicit list", d["error"])

    def test_apply_reads_the_review_saida_and_never_recomputes_it(self):
        d = sim({**ARGS, "modo": "aplicar", "aplicar": ["c4-01"]})
        self.assertIn("aplicar needs saida", d["error"] or "")
        self.assertEqual(d["calls"], [])
        d = sim({**ARGS, "modo": "aplicar", "saida": "/t/rev-antiga", "aplicar": ["c4-01"]})
        self.assertIsNone(d["error"])
        self.assertIn("/t/rev-antiga/revisao.md", d["calls"][0]["prompt"])
        self.assertEqual(d["result"]["saida"], "/t/rev-antiga")


class LensSections(unittest.TestCase):
    def test_script_lenses_match_lenses_md(self):
        with open(SCRIPT, encoding="utf-8") as fh:
            block = fh.read().split("const LENSES = [", 1)[1].split("]", 1)[0]
        keys = re.findall(r"key: '([a-z-]+)'", block)
        self.assertEqual(len(keys), 10)
        with open(os.path.join(REF, "lenses.md"), encoding="utf-8") as fh:
            sections = re.findall(r"^## ([a-z-]+)\s*$", fh.read(), re.M)
        self.assertEqual(sorted(sections), sorted(keys + ["ficha", "critico", "refutador"]))

    def test_meta_is_a_literal_and_script_avoids_the_clock(self):
        with open(SCRIPT, encoding="utf-8") as fh:
            src = fh.read()
        meta = src.split("export const meta = {", 1)[1].split("\n}\n", 1)[0]
        self.assertNotIn("${", meta)
        self.assertNotRegex(src, r"Date\.now\(\)|Math\.random\(\)|new Date\(\)")


class Helpers(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp)

    def write(self, name, text):
        path = os.path.join(self.tmp, name)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(text)
        return path

    def test_scout_accepts_numbered_markdown_and_refuses_unnumbered(self):
        good = self.write("c.md", "# Contrato\n\n1. Objeto\n\n1.1 O objeto é X.\n\n2. Preço\n")
        out = helper("docx-scout.py", good, "--out", os.path.join(self.tmp, "s"))
        self.assertEqual(out.returncode, 0, out.stderr)
        self.assertIn(f"contrato: {good}", out.stdout)
        bad = self.write("d.md", "Contrato sem numeração.\n\nOutro parágrafo.\n")
        self.assertEqual(helper("docx-scout.py", bad, "--out", self.tmp).returncode, 3)

    def test_scout_rebuilds_docx_numbering_from_numpr(self):
        try:
            import docx
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError:
            self.skipTest("python-docx not installed")
        doc = docx.Document()
        # A decimal list of our own, so the test does not depend on the template's lists.
        root = doc.part.numbering_part.element
        an = OxmlElement("w:abstractNum"); an.set(qn("w:abstractNumId"), "90")
        for lvl in ("0", "1"):
            el = OxmlElement("w:lvl"); el.set(qn("w:ilvl"), lvl)
            fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "decimal"); el.append(fmt)
            an.append(el)
        root.insert(0, an)
        num = OxmlElement("w:num"); num.set(qn("w:numId"), "90")
        ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), "90"); num.append(ref)
        root.append(num)
        numid = "90"
        for text, lvl in (("Objeto", 0), ("O objeto é X.", 1), ("Preço", 0)):
            p = doc.add_paragraph(text)
            numpr = OxmlElement("w:numPr")
            il, ni = OxmlElement("w:ilvl"), OxmlElement("w:numId")
            il.set(qn("w:val"), str(lvl)); ni.set(qn("w:val"), numid)
            numpr.append(il); numpr.append(ni)
            p._p.get_or_add_pPr().append(numpr)
        src = os.path.join(self.tmp, "c.docx")
        doc.save(src)
        out = helper("docx-scout.py", src, "--out", os.path.join(self.tmp, "s"))
        md = os.path.join(self.tmp, "s", "contrato.md")
        self.assertEqual(out.returncode, 0, out.stderr)
        with open(md, encoding="utf-8") as fh:
            text = fh.read()
        self.assertIn("1. Objeto", text)
        self.assertIn("1.1. O objeto é X.", text)
        self.assertIn("2. Preço", text)
        # A file the script did not write is never overwritten.
        with open(md, "w", encoding="utf-8") as fh:
            fh.write("nota do usuário\n")
        self.assertEqual(helper("docx-scout.py", src, "--out", os.path.join(self.tmp, "s")).returncode, 2)

    def test_scout_stops_when_style_numbering_is_lost(self):
        try:
            import docx
            from docx.enum.style import WD_STYLE_TYPE
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError:
            self.skipTest("python-docx not installed")
        doc = docx.Document()
        root = doc.part.numbering_part.element
        an = OxmlElement("w:abstractNum"); an.set(qn("w:abstractNumId"), "91")
        el = OxmlElement("w:lvl"); el.set(qn("w:ilvl"), "0")
        fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "decimal"); el.append(fmt); an.append(el)
        root.insert(0, an)
        num = OxmlElement("w:num"); num.set(qn("w:numId"), "91")
        ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), "91"); num.append(ref)
        root.append(num)
        # Word's usual clause list: the numPr lives on the paragraph STYLE, not on the paragraph.
        style = doc.styles.add_style("Clausula", WD_STYLE_TYPE.PARAGRAPH)
        numpr = OxmlElement("w:numPr")
        il, ni = OxmlElement("w:ilvl"), OxmlElement("w:numId")
        il.set(qn("w:val"), "0"); ni.set(qn("w:val"), "91")
        numpr.append(il); numpr.append(ni)
        style.element.get_or_add_pPr().append(numpr)
        doc.add_paragraph("1. Preâmbulo digitado à mão")
        for text in ("Objeto", "Preço", "Prazo"):
            doc.add_paragraph(text, style="Clausula")
        src = os.path.join(self.tmp, "c.docx")
        doc.save(src)
        out = helper("docx-scout.py", src, "--out", os.path.join(self.tmp, "s"))
        self.assertEqual(out.returncode, 3, out.stdout + out.stderr)
        self.assertIn("1 parágrafos numerados, 3 com w:numPr", out.stderr)

    def test_brutos_assembles_with_workflow_ids(self):
        saida = os.path.join(self.tmp, "s")
        self.write("s/brutos/tributaria.json", json.dumps({"resumo": "r", "findings": [
            {"clausula": "5.4", "severidade": "alta", "titulo": "t", "problema": "p", "proposta": "q"}]}))
        out = helper("brutos.py", saida)
        self.assertEqual(out.returncode, 0, out.stderr)
        with open(os.path.join(saida, "achados-brutos.json"), encoding="utf-8") as fh:
            self.assertEqual(json.load(fh)["achados"][0]["id"], "tributaria-1")
        self.assertTrue(os.path.isfile(os.path.join(saida, "achados-brutos.md")))
        self.assertEqual(helper("brutos.py", os.path.join(self.tmp, "vazio")).returncode, 2)
        # Its own files it rewrites.
        self.assertEqual(helper("brutos.py", saida).returncode, 0)

    def test_brutos_never_overwrites_a_json_it_did_not_write(self):
        saida = os.path.join(self.tmp, "s")
        self.write("s/brutos/tributaria.json", json.dumps({"resumo": "r", "findings": []}))
        foreign = self.write("s/achados-brutos.json", '{"nota": "do usuário"}')
        out = helper("brutos.py", saida)
        self.assertEqual(out.returncode, 2)
        self.assertIn("não foi escrito por este script", out.stderr)
        with open(foreign, encoding="utf-8") as fh:
            self.assertEqual(fh.read(), '{"nota": "do usuário"}')
        self.assertFalse(os.path.exists(os.path.join(saida, "achados-brutos.md")))

    def test_ids_lists_the_report_altas(self):
        rep = self.write("s/revisao.md", "## 2. Corrigir\n\n### c4-01 · cláusula 4.2\n\nx\n\n### anexo-iv-01 · cláusula Anexo IV\n")
        out = helper("apply-edits.py", "ids", rep)
        self.assertEqual(out.stdout.split(), ["c4-01", "anexo-iv-01"])

    def test_apply_snapshots_then_writes_within_budget(self):
        c = self.write("c.md", "1. O preço é de dez reais.\n\n2. Prazo de doze meses.\n")
        self.write("s/aplicar/c1.json", json.dumps({"edits": [{"id": "c1-01", "antes": "dez reais", "depois": "onze reais"}]}))
        out = helper("apply-edits.py", "apply", "--contrato", c, "--edits", os.path.join(self.tmp, "s", "aplicar"), "--data", "2026-09-21")
        self.assertEqual(out.returncode, 0, out.stderr)
        with open(os.path.join(self.tmp, "versoes", "c-2026-09-21.md"), encoding="utf-8") as fh:
            self.assertIn("dez reais", fh.read())
        with open(c, encoding="utf-8") as fh:
            self.assertIn("onze reais", fh.read())

    def test_apply_overrun_writes_nothing(self):
        text = "1. O preço é de dez reais.\n"
        c = self.write("c.md", text)
        self.write("s/aplicar/c1.json", json.dumps({"edits": [{"id": "c1-01", "antes": "dez reais", "depois": "dez reais corrigidos anualmente pelo índice oficial"}]}))
        out = helper("apply-edits.py", "apply", "--contrato", c, "--edits", os.path.join(self.tmp, "s", "aplicar"), "--data", "2026-09-21", "--folga", "+10%")
        self.assertEqual(out.returncode, 4)
        self.assertIn("ESTOURO", out.stderr)
        with open(c, encoding="utf-8") as fh:
            self.assertEqual(fh.read(), text)
        self.assertFalse(os.path.exists(os.path.join(self.tmp, "versoes")))

    def test_apply_refuses_a_folga_without_percent(self):
        text = "1. O preço é de dez reais.\n"
        c = self.write("c.md", text)
        self.write("s/aplicar/c1.json", json.dumps({"edits": [{"id": "c1-01", "antes": "dez reais", "depois": "dez reais corrigidos anualmente"}]}))
        out = helper("apply-edits.py", "apply", "--contrato", c, "--edits", os.path.join(self.tmp, "s", "aplicar"), "--data", "2026-09-21", "--folga", "10")
        self.assertEqual(out.returncode, 2)
        self.assertIn("precisa de '%'", out.stderr)
        with open(c, encoding="utf-8") as fh:
            self.assertEqual(fh.read(), text)
        self.assertIn("a bare number is refused", helper("apply-edits.py", "apply", "-h").stdout)

    def _docx_apply(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx not installed")
        doc = docx.Document(); doc.add_paragraph("1. O preço é de dez reais.")
        orig = os.path.join(self.tmp, "c.docx"); doc.save(orig)
        md = self.write("s/contrato.md", "1. O preço é de dez reais.\n")
        self.write("s/aplicar/c1.json", json.dumps({"edits": [{"id": "c1-01", "antes": "dez reais", "depois": "onze reais"}]}))
        argv = ("apply", "--contrato", md, "--original", orig, "--edits", os.path.join(self.tmp, "s", "aplicar"), "--data", "2026-09-21")
        return md, argv

    def test_apply_on_a_docx_snapshots_the_text_as_markdown(self):
        md, argv = self._docx_apply()
        out = helper("apply-edits.py", *argv)
        self.assertEqual(out.returncode, 0, out.stderr)
        with open(os.path.join(self.tmp, "versoes", "c-2026-09-21.md"), encoding="utf-8") as fh:
            self.assertIn("dez reais", fh.read())
        self.assertTrue(os.path.isfile(os.path.join(self.tmp, "versoes", "c-2026-09-21.docx")))
        with open(md, encoding="utf-8") as fh:
            self.assertIn("onze reais", fh.read())

    def test_apply_never_overwrites_a_docx_it_did_not_write(self):
        md, argv = self._docx_apply()
        import docx
        dest = os.path.join(self.tmp, "c-aplicado-2026-09-21.docx")
        foreign = docx.Document(); foreign.add_paragraph("versão do advogado"); foreign.save(dest)
        out = helper("apply-edits.py", *argv)
        self.assertEqual(out.returncode, 5)
        self.assertIn("não foi escrito por este script", out.stderr)
        self.assertEqual(docx.Document(dest).paragraphs[0].text, "versão do advogado")
        with open(md, encoding="utf-8") as fh:
            self.assertIn("dez reais", fh.read())
        self.assertFalse(os.path.exists(os.path.join(self.tmp, "versoes")))
        # What this script writes carries the mark it checks.
        os.remove(dest)
        self.assertEqual(helper("apply-edits.py", *argv).returncode, 0)
        self.assertEqual(docx.Document(dest).core_properties.comments, "revisao-contrato apply-edits.py")

    def test_apply_refuses_an_ambiguous_anchor(self):
        c = self.write("c.md", "1. dez e dez.\n")
        self.write("s/aplicar/c1.json", json.dumps({"edits": [{"id": "c1-01", "antes": "dez", "depois": "onze"}]}))
        out = helper("apply-edits.py", "apply", "--contrato", c, "--edits", os.path.join(self.tmp, "s", "aplicar"), "--data", "2026-09-21")
        self.assertEqual(out.returncode, 2)
        self.assertIn("2 vezes", out.stderr)


if __name__ == "__main__":
    unittest.main()
