#!/usr/bin/env python3
"""Apply step of revisao-contrato: list the report's ids, then apply chosen edits under a word budget.

    python3 apply-edits.py ids <saida>/revisao.md
    python3 apply-edits.py apply --contrato <text.md> --original <file as received> \
        --edits <saida>/aplicar --data YYYY-MM-DD [--folga +10%]

`ids` prints every alta id of the report, one per line — what `aplicar: "todas"` means.
`apply` reads every `<edits>/*.json` ({"edits": [{"id", "antes", "depois"}]}), applies them in
order to the text, and measures words before and after. Over `before x (1 + folga)` it writes
NOTHING, prints the overrun and each finding's word delta, and exits 4. `--folga` is a
percentage and must carry `%` (`+10%`); a bare number is refused, since `10` would read as 1000%.
Otherwise it snapshots the current text to `<original dir>/versoes/<stem>-<data>.md` first — and,
for a `.docx` original, the `.docx` itself to `versoes/<stem>-<data>.docx` — then writes the
text; for a `.docx` original it also writes `<stem>-aplicado-<data>.docx` beside it, marked as
its own in the document's comments property.
Exit 2: usage or read error, or an `antes` that does not match the text exactly once.
Exit 5: a snapshot path already holds a different file, or `<stem>-aplicado-<data>.docx` exists
and this script did not write it.
"""

import argparse
import filecmp
import glob
import json
import os
import re
import shutil
import sys

HEADING = re.compile(r"^###\s+(\S+)\s+·\s+cl[aá]usula\b", re.M)
DOCX_MARK = "revisao-contrato apply-edits.py"


def read_text(path):
    if not os.path.isfile(path):
        raise OSError(f"{path} não é um arquivo regular")
    with open(path, encoding="utf-8-sig") as fh:
        return fh.read().replace("﻿", "")


def words(text):
    return len(text.split())


def parse_folga(s):
    s = s.strip().lstrip("+")
    if not s.endswith("%"):
        raise ValueError(f"--folga precisa de '%' (ex.: +10%), recebeu {s!r}")
    return float(s[:-1]) / 100


def cmd_ids(opts):
    text = read_text(opts.report)
    ids = list(dict.fromkeys(HEADING.findall(text)))
    if not ids:
        print(f"nenhum bloco '### <id> · cláusula' em {opts.report}", file=sys.stderr)
        return 2
    print("\n".join(ids))
    return 0


def write_docx(md_text, dest):
    import docx  # python-docx
    doc = docx.Document()
    doc.core_properties.comments = DOCX_MARK
    for line in md_text.splitlines():
        if line.startswith("<!--"):
            continue
        m = re.match(r"^(#+)\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 9))
        else:
            doc.add_paragraph(line)
    doc.save(dest)


def docx_is_ours(path):
    import docx  # python-docx
    try:
        return docx.Document(path).core_properties.comments == DOCX_MARK
    except Exception:  # not a readable .docx: certainly not one this script wrote
        return False


def cmd_apply(opts):
    folga = parse_folga(opts.folga)
    text = read_text(opts.contrato)
    edits = []
    for path in sorted(glob.glob(os.path.join(opts.edits, "*.json"))):
        data = json.loads(read_text(path))
        edits.extend(data.get("edits", []))
    if not edits:
        print(f"erro: nenhuma edição em {opts.edits}/*.json", file=sys.stderr)
        return 2
    before = words(text)
    deltas = []
    for e in edits:
        antes, depois = e.get("antes", ""), e.get("depois", "")
        hits = text.count(antes) if antes else 0
        if hits != 1:
            print(f"erro: {e.get('id')}: o trecho 'antes' aparece {hits} vezes no texto (precisa ser 1)", file=sys.stderr)
            return 2
        text = text.replace(antes, depois, 1)
        deltas.append((e.get("id"), words(depois) - words(antes)))
    after = words(text)
    limit = int(before * (1 + folga))
    for fid, d in deltas:
        print(f"{fid}: {d:+d} palavras")
    print(f"palavras: {before} → {after} (limite {limit}, folga {folga:.0%})")
    if after > limit:
        print(f"ESTOURO: {after - limit} palavras acima do limite; nada foi escrito. "
              f"Compense com cortes do plano de clareza ou aplique menos ids.", file=sys.stderr)
        return 4
    orig = opts.original or opts.contrato
    stem, ext = os.path.splitext(os.path.basename(orig))
    odir = os.path.dirname(os.path.abspath(orig))
    vdir = os.path.join(odir, "versoes")
    # The current TEXT is always snapshotted as .md; a non-Markdown original is kept as well.
    snaps = [(opts.contrato, os.path.join(vdir, f"{stem}-{opts.data}.md"))]
    if ext.lower() != ".md":
        snaps.append((orig, os.path.join(vdir, f"{stem}-{opts.data}{ext}")))
    dest = os.path.join(odir, f"{stem}-aplicado-{opts.data}.docx") if ext.lower() == ".docx" else None
    # Every refusal comes before the first write, a missing python-docx included.
    if dest:
        import docx  # noqa: F401
    for src, snap in snaps:
        if os.path.exists(snap) and not filecmp.cmp(snap, src, shallow=False):
            print(f"erro: {snap} já existe com outro conteúdo; mova-o antes de aplicar", file=sys.stderr)
            return 5
    if dest and os.path.exists(dest) and not docx_is_ours(dest):
        print(f"erro: {dest} existe e não foi escrito por este script; mova-o antes de aplicar", file=sys.stderr)
        return 5
    os.makedirs(vdir, exist_ok=True)
    for src, snap in snaps:
        if not os.path.exists(snap):
            shutil.copy2(src, snap)
        print(f"versão anterior: {snap}")
    with open(opts.contrato, "w", encoding="utf-8") as fh:
        fh.write(text)
    print(f"texto aplicado: {opts.contrato}")
    if dest:
        write_docx(text, dest)
        print(f"docx regenerado: {dest}")
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    p_ids = sub.add_parser("ids")
    p_ids.add_argument("report")
    p_apply = sub.add_parser("apply")
    p_apply.add_argument("--contrato", required=True, help="the Markdown text the review read")
    p_apply.add_argument("--original", help="the file as received, when it is not --contrato (e.g. the .docx)")
    p_apply.add_argument("--edits", required=True)
    p_apply.add_argument("--data", required=True)
    p_apply.add_argument("--folga", default="+10%",
                         help="allowed growth of the text, a percentage with %% (default +10%%); a bare number is refused")
    opts = ap.parse_args(argv)
    if opts.cmd == "apply" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", opts.data):
        print("erro: --data deve ser YYYY-MM-DD", file=sys.stderr)
        return 2
    try:
        return cmd_ids(opts) if opts.cmd == "ids" else cmd_apply(opts)
    except (OSError, UnicodeDecodeError, ValueError) as exc:
        print(f"erro: {exc}", file=sys.stderr)
        return 2
    except ImportError:
        print("erro: python-docx ausente (pip install python-docx)", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
